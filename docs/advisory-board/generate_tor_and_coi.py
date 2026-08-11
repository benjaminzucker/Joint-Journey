#!/usr/bin/env python3
"""Generate formatted Word documents for Terms of Reference and CoI Register."""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))

# ── Brand colours ──
GREEN_DARK = RGBColor(0x47, 0x59, 0x53)
GREEN_MID = RGBColor(0x2F, 0x85, 0x5A)
GREEN_LIGHT_BG = RGBColor(0xEE, 0xF3, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x66, 0x66, 0x66)


def set_cell_shading(cell, color_hex):
    """Set background shading on a table cell."""
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): color_hex,
        qn('w:val'): 'clear',
    })
    shading.append(shd)


def style_doc(doc):
    """Apply base styles to the document."""
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = BLACK
    pf = style.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)


def add_logo_header(doc, title, subtitle=None):
    """Add a branded header block."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run('Joint')
    run.font.size = Pt(22)
    run.font.color.rgb = GREY
    run.font.name = 'Calibri'
    run = p.add_run('Journey')
    run.font.size = Pt(22)
    run.font.color.rgb = GREEN_DARK
    run.font.bold = True
    run.font.name = 'Calibri'

    # Title
    h = doc.add_heading(title, level=1)
    for run in h.runs:
        run.font.color.rgb = GREEN_DARK
        run.font.size = Pt(18)

    if subtitle:
        p = doc.add_paragraph()
        run = p.add_run(subtitle)
        run.font.size = Pt(10)
        run.font.color.rgb = GREY
        run.font.italic = True

    # Thin green line
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run('─' * 80)
    run.font.size = Pt(6)
    run.font.color.rgb = GREEN_MID


def add_section_heading(doc, text):
    """Add a styled section heading."""
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.color.rgb = GREEN_DARK
        run.font.size = Pt(14)


def add_table(doc, headers, rows, col_widths=None):
    """Add a styled table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = WHITE
        run.font.name = 'Calibri'
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_cell_shading(cell, '475953')

    # Data rows
    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(10)
            run.font.name = 'Calibri'
            if r_idx % 2 == 1:
                set_cell_shading(cell, 'EEF3EC')

    # Column widths
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    return table


def add_bullet(doc, text, bold_prefix=None):
    """Add a bullet point, optionally with a bold prefix."""
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.font.bold = True
        run.font.size = Pt(11)
        p.add_run(text)
    else:
        p.add_run(text)


def add_note_box(doc, text):
    """Add a highlighted note/callout paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run('⚠️  ' + text)
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x6B, 0x4A, 0x00)


# ═══════════════════════════════════════════════════════════════
# DOCUMENT 1: TERMS OF REFERENCE
# ═══════════════════════════════════════════════════════════════
def generate_tor():
    doc = Document()
    style_doc(doc)

    add_logo_header(
        doc,
        'Clinical Advisory Board\nTerms of Reference',
        'Version 0.1 (draft)  ·  Owner: Mr Benjamin Zucker, Founder & CSO'
    )

    add_note_box(doc, 'Working draft. Review with the board at its first meeting and adopt formally.')

    # ── 1. Purpose ──
    add_section_heading(doc, '1. Purpose')
    doc.add_paragraph(
        'The Clinical Advisory Board ("the Board") provides independent, multidisciplinary '
        'clinical input to Joint Journey, a digital prehabilitation programme for adults '
        'awaiting hip or knee replacement. The Board exists to:'
    )
    add_bullet(doc, ' — ensure the exercise, nutrition and mental-health/pain-preparation '
               'content is safe, evidence-based and appropriate.', 'Assure clinical content')
    add_bullet(doc, ' — contribute to hazard identification and mitigation '
               '(feeding the DCB0129 clinical safety case).', 'Support clinical risk management')
    add_bullet(doc, ' — advise on what can and cannot be claimed for the product.', 'Guide claims')
    add_bullet(doc, ' — provide named clinical oversight for users, NHS trusts and funders.', 'Lend credibility')

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run('The Board is advisory')
    run.font.bold = True
    p.add_run(': it informs decisions but does not direct the company. Final decisions rest '
              'with the company/founder, except that ')
    run = p.add_run('clinical content must not be published without the relevant advisor\'s '
                    'documented sign-off.')
    run.font.bold = True

    # ── 2. Membership ──
    add_section_heading(doc, '2. Membership')
    add_table(doc,
              ['Seat', 'Discipline', 'Holder'],
              [
                  ['Chair', 'Founder / surgical oversight (orthopaedics)', 'Mr Benjamin Zucker'],
                  ['Member', 'MSK / orthopaedic Physiotherapist (HCPC)', ''],
                  ['Member', 'Clinical / Health Psychologist (HCPC) or CBT therapist (BABCP)', ''],
                  ['Member (recommended)', 'Registered Dietitian (HCPC)', ''],
                  ['Optional', 'Patient representative (PPIE) — lived experience', ''],
                  ['Optional', 'Consultant Orthopaedic Surgeon', ''],
                  ['Optional', 'GP / pre-operative nurse', ''],
              ],
              col_widths=[1.5, 3.5, 1.8])

    doc.add_paragraph()
    add_bullet(doc, 'All clinical members must hold current professional registration (HCPC/GMC/BABCP).')
    add_bullet(doc, 'Members serve a renewable term of 2 years.')
    add_bullet(doc, 'The Clinical Safety Officer (CSO) for the product is Mr Benjamin Zucker. '
               'The Board supports, but does not replace, the CSO.')

    # ── 3. Roles & responsibilities ──
    add_section_heading(doc, '3. Roles & Responsibilities')
    p = doc.add_paragraph()
    run = p.add_run('Members will:')
    run.font.bold = True
    add_bullet(doc, 'Review the content for their pillar (initial and on material change) and '
               'record sign-off in the Content Review and Sign-off Log.')
    add_bullet(doc, 'Help identify hazards and mitigations in their domain (Hazard Log).')
    add_bullet(doc, 'Declare conflicts of interest (CoI Register).')
    add_bullet(doc, 'Attend meetings (or send comments) and act within their professional scope.')

    p = doc.add_paragraph()
    run = p.add_run('Members will NOT:')
    run.font.bold = True
    add_bullet(doc, 'Take responsibility for company decisions or operations.')
    add_bullet(doc, 'Provide individual clinical care to users through this role.')

    # ── 4. Independence & CoI ──
    add_section_heading(doc, '4. Independence & Conflicts of Interest')
    doc.add_paragraph(
        'Members should be independent of any NHS organisation that may later commission or '
        'purchase Joint Journey, to avoid procurement conflicts. Where a member has such a '
        'connection, it must be declared and managed (see CoI Policy).'
    )
    doc.add_paragraph(
        'Conflicts are declared on joining, recorded in the register, and revisited at every meeting.'
    )

    # ── 5. Meetings ──
    add_section_heading(doc, '5. Meetings')
    add_bullet(doc, ' Quarterly (≈4/year), plus async reviews as content is developed.', 'Frequency:')
    add_bullet(doc, ' The Chair plus at least 2 members; the relevant pillar\'s advisor must be '
               'present (or have commented) for any decision on that pillar.', 'Quorum:')
    add_bullet(doc, ' Taken for every meeting, recording attendance, decisions, sign-offs, '
               'actions and declared interests.', 'Minutes:')
    add_bullet(doc, ' By consensus; the Chair holds a casting view on non-clinical matters. '
               'Clinical content sign-off rests with the relevant registered advisor.', 'Decisions:')

    # ── 6. Compensation ──
    add_section_heading(doc, '6. Compensation')
    doc.add_paragraph(
        'As set out in each member\'s Advisor Agreement. Default approach: honoraria and/or '
        'pro-bono initially; equity (FAST-style vesting) once the company is incorporated. '
        'Confirm per member.'
    )

    # ── 7. Confidentiality & IP ──
    add_section_heading(doc, '7. Confidentiality & IP')
    add_bullet(doc, 'Members keep company information confidential (see Advisor Agreement).')
    add_bullet(doc, 'Intellectual property in content/advice created for the company is '
               'assigned to the company (see Advisor Agreement).')

    # ── 8. Indemnity ──
    add_section_heading(doc, '8. Indemnity & Responsibility')
    doc.add_paragraph(
        'Members provide advice in good faith within their professional competence. '
        'The company (via the CSO) holds responsibility for the product and its deployment. '
        'Members should confirm their advisory activity is covered by their professional '
        'indemnity, or rely on the company\'s arrangements as agreed.'
    )

    # ── 9. Review ──
    add_section_heading(doc, '9. Review')
    doc.add_paragraph(
        'This Terms of Reference is reviewed annually and updated as the product and regulatory '
        'position evolve (e.g. at MHRA/medical-device decision points).'
    )

    # ── Signature block ──
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run('─' * 80)
    run.font.size = Pt(6)
    run.font.color.rgb = GREEN_MID

    p = doc.add_paragraph()
    run = p.add_run('Adopted by the Board on: ')
    run.font.bold = True
    p.add_run('___________________')

    p = doc.add_paragraph()
    run = p.add_run('Chair signature: ')
    run.font.bold = True
    p.add_run('___________________')

    p = doc.add_paragraph()
    run = p.add_run('Date: ')
    run.font.bold = True
    p.add_run('___________________')

    path = os.path.join(OUTPUT_DIR, 'Terms of Reference.docx')
    doc.save(path)
    print(f'✅ Saved: {path}')


# ═══════════════════════════════════════════════════════════════
# DOCUMENT 2: CONFLICT OF INTEREST REGISTER
# ═══════════════════════════════════════════════════════════════
def generate_coi():
    doc = Document()
    style_doc(doc)

    add_logo_header(
        doc,
        'Conflict of Interest Register',
        'Master record of declared interests  ·  Reviewed at least annually and at each Board meeting'
    )

    doc.add_paragraph(
        'This register records all interests declared by Clinical Advisory Board members, '
        'the founder, and anyone advising or working on Joint Journey. '
        'See the Conflict of Interest Policy for definitions and management procedures.'
    )

    add_note_box(doc, '"Nil declared" is a valid and useful entry — record it explicitly.')

    # ── Register table ──
    add_section_heading(doc, 'Declared Interests')

    headers = ['#', 'Name', 'Role / Seat', 'Date Declared', 'Interest Declared',
               'Type', 'How Managed', 'Reviewed']
    rows = [
        ['1', 'Mr Benjamin Zucker', 'Founder / Chair / CSO', '', 
         'Founder & (intended) shareholder of Joint Journey', 'Financial',
         'Inherent to role; noted', ''],
        ['2', '', 'Consultant Orthopaedic Surgeon', '', '', '', '', ''],
        ['3', '', 'Physiotherapy Advisor', '', '', '', '', ''],
        ['4', '', 'Psychology Advisor', '', '', '', '', ''],
        ['5', '', 'Dietetics Advisor', '', '', '', '', ''],
        ['6', '', 'Patient Representative', '', '', '', '', ''],
    ]

    add_table(doc, headers, rows, col_widths=[0.3, 1.1, 1.1, 0.7, 1.5, 0.7, 1.0, 0.7])

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    run = p.add_run('Instructions: ')
    run.font.bold = True
    p.add_run('Add a new row whenever a member joins or an interest changes. '
              'Types: Financial / Personal / Loyalty / Indirect.')

    # ── CoI Policy Summary ──
    add_section_heading(doc, 'Conflict of Interest Policy — Summary')

    doc.add_paragraph(
        'This policy identifies, declares and manages any interests that could (or could '
        'appear to) improperly influence the advice given to, or decisions made by, Joint Journey.'
    )

    p = doc.add_paragraph()
    run = p.add_run('What is a conflict of interest?')
    run.font.bold = True
    doc.add_paragraph(
        'An interest that might affect, or be seen to affect, a person\'s objectivity.'
    )

    p = doc.add_paragraph()
    run = p.add_run('Types:')
    run.font.bold = True
    add_bullet(doc, ' — payments, equity, shares, consultancy with competitors or suppliers.', 'Financial')
    add_bullet(doc, ' — friendships, family relationships, reputation.', 'Non-financial / Personal')
    add_bullet(doc, ' — employment or office at an NHS organisation that may commission or buy the product; '
               'roles with competing products.', 'Loyalty / Role')
    add_bullet(doc, ' — interests of a close family member or close associate.', 'Indirect')

    p = doc.add_paragraph()
    run = p.add_run('Duties — everyone covered by this policy must:')
    run.font.bold = True
    add_bullet(doc, 'Declare relevant interests on joining, in this register.')
    add_bullet(doc, 'Update the register promptly when interests change.')
    add_bullet(doc, 'Disclose at the start of each meeting any interest relevant to the agenda.')
    add_bullet(doc, 'Step back from any decision where they have a material conflict '
               '(the minutes must record this).')

    p = doc.add_paragraph()
    run = p.add_run('Managing a conflict — ')
    run.font.bold = True
    p.add_run('depending on severity, the Board may: record and note it; exclude the person '
              'from the relevant discussion/decision; remove the relevant item from their remit; '
              'or in serious cases, end the advisory relationship. The Chair decides how each '
              'conflict is managed and records it in the minutes and register.')

    # ── Review & signature ──
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run('─' * 80)
    run.font.size = Pt(6)
    run.font.color.rgb = GREEN_MID

    p = doc.add_paragraph()
    run = p.add_run('Last full review: ')
    run.font.bold = True
    p.add_run('___________________')

    p = doc.add_paragraph()
    run = p.add_run('Reviewed by (Chair): ')
    run.font.bold = True
    p.add_run('___________________')

    p = doc.add_paragraph()
    run = p.add_run('Date: ')
    run.font.bold = True
    p.add_run('___________________')

    path = os.path.join(OUTPUT_DIR, 'Conflict of Interest Register.docx')
    doc.save(path)
    print(f'✅ Saved: {path}')


if __name__ == '__main__':
    generate_tor()
    generate_coi()
    print('\nDone! Both documents generated.')
